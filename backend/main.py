"""Parent Account Finder — backend (Substrait upload mode, OceanBase/MySQL-wire).

On-demand companion to the sibling Account Parent-Mapping Monitor app: instead of a
full dashboard of every duplicate-name group, Sales Ops pastes or uploads specific
Account IDs / Opportunity IDs and gets back just that Account's likely duplicate-name
matches (exact name, fuzzy name, and shared phone/NPWP/identification-number), each
with its current and a suggested Parent Account — for manual review and correction
directly in CRM. Read-only — never writes to CRM.

Like the Monitor app, this mirrors CRM Accounts (Indonesia record type only) into a
local table rather than querying CRM live on every lookup: exact-name and identity
matching need to scan the *whole* Indonesia Account set, not just what a live search
call happens to rank highly. The mirror is kept in sync by the same
weekly-background-job + "Sync now" pattern as the Monitor app.
"""
import asyncio
import csv
import io
import json
import os
import re
from collections import defaultdict
from contextlib import asynccontextmanager
from datetime import datetime
from urllib.parse import unquote, urlparse

import asyncmy
import openpyxl
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from rapidfuzz import fuzz

import crm_client as crm

_pool = None
_sync_task = None
_sync_in_progress = False

SYNC_INTERVAL_SECONDS = int(os.getenv("SYNC_INTERVAL_SECONDS", str(7 * 24 * 3600)))  # weekly

# This CRM tenant spans multiple countries; this app is scoped to Indonesia only —
# filtered server-side in crm_client by record_type_id (the CRM's own Record Type
# foreign key, NOT the free-text account_record_type field — see the sibling Monitor
# app's crm_client.py docstring for why). 10 confirmed live.
ACCOUNT_RECORD_TYPE_ID = 10

# Fuzzy matching: catches near-duplicate names (typos, extra suffix words) that exact
# normalization misses. token_sort_ratio, 0-100 scale — 88 is a fairly tight bar,
# chosen (same as the Monitor app) to keep false positives low.
FUZZY_SIMILARITY_THRESHOLD = 88
# 125k+ Accounts means ~100k+ distinct names — full pairwise comparison is O(n^2) and
# infeasible. Blocking by first (non-generic) word keeps comparisons to within
# same-first-word buckets; a bucket bigger than this is skipped entirely.
FUZZY_MAX_BLOCK_SIZE = 300
_FUZZY_BLOCK_STOPWORDS = {"pt", "cv", "ud", "toko", "the", "pt.", "cv."}

# Identity matching: Accounts sharing the same phone / NPWP (tax_id) / identification
# number almost certainly belong to the same company, even when names differ too much
# for name-based matching. (digits_column, raw_column, display label) — the digits
# column is computed at sync time (see do_sync) so a lookup is an indexed equality
# query instead of a full-table scan-and-normalize.
IDENTITY_FIELDS = [
    ("phone_digits", "phone", "Phone"),
    ("tax_id_digits", "tax_id", "NPWP"),
    ("identification_number_digits", "identification_number", "Identification Number"),
]
_IDENTITY_MIN_DIGITS = 6  # below this, too short to be a real number — likely a stray "0" or similar

BATCH_MAX_IDS = 1000  # sanity cap on one paste/upload run

_MIRROR_FIELDS = [
    "id", "name", "normalized_name", "owner_id", "owner_name", "parent_account_id",
    "phone", "tax_id", "identification_number", "created_at_crm",
]


def _row_to_dict(row) -> dict:
    return dict(zip(_MIRROR_FIELDS, row))


def _normalize_identity_value(value: str | None) -> str | None:
    if not value:
        return None
    digits = re.sub(r"\D", "", value)
    return digits if len(digits) >= _IDENTITY_MIN_DIGITS else None


def _dsn() -> dict:
    u = urlparse(os.environ["DATABASE_URL"])
    return {
        "host": u.hostname,
        "port": u.port or 2881,
        "user": unquote(u.username or ""),
        "password": unquote(u.password or ""),
        "db": (u.path or "/").lstrip("/"),
    }


async def _run_sync(trigger: str) -> None:
    """Runs do_sync in the background and clears the in-progress flag when done. A
    full paginated Account pull can take well over a minute — long enough to hit the
    platform ingress's request timeout — so callers must never await this directly
    from an HTTP handler."""
    global _sync_in_progress
    try:
        await do_sync(trigger=trigger)
    except Exception:
        pass  # do_sync already records the error on the sync_runs row
    finally:
        _sync_in_progress = False


async def _sync_loop():
    global _sync_in_progress
    while True:
        await asyncio.sleep(SYNC_INTERVAL_SECONDS)
        if _sync_in_progress:
            continue  # a manual sync is already running; skip this tick
        _sync_in_progress = True
        await _run_sync("scheduled")


@asynccontextmanager
async def lifespan(app: FastAPI):
    global _pool, _sync_task
    if os.getenv("DATABASE_URL"):
        _pool = await asyncmy.create_pool(**_dsn(), autocommit=True)
    _sync_task = asyncio.create_task(_sync_loop())
    yield
    if _sync_task:
        _sync_task.cancel()
    if _pool is not None:
        _pool.close()
        await _pool.wait_closed()


app = FastAPI(title="Parent Account Finder", lifespan=lifespan)

_PUNCT_RE = re.compile(r"[^\w\s]", re.UNICODE)
_WS_RE = re.compile(r"\s+")


def normalize_name(name: str) -> str:
    """lowercase, strip punctuation, collapse whitespace — matches the literal
    exact-duplicate pattern seen in real Account data (same name string reused
    verbatim for related shops/branches). Same logic as the sibling Monitor app."""
    if not name:
        return ""
    stripped = _PUNCT_RE.sub(" ", name.lower())
    return _WS_RE.sub(" ", stripped).strip()


def _parse_dt(value) -> datetime | None:
    """Parses a CRM timestamp into a naive UTC datetime — MySQL DATETIME columns are
    naive, and mixing naive/aware datetimes later (e.g. in min()) raises TypeError."""
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None
    return parsed.replace(tzinfo=None)


def _fuzzy_block_key(normalized_name: str) -> str:
    words = normalized_name.split()
    for w in words:
        if w not in _FUZZY_BLOCK_STOPWORDS and len(w) >= 3:
            return w
    return words[0] if words else normalized_name


def _compute_fuzzy_matches(distinct_names: list[str]) -> list[tuple[str, str, int]]:
    """Pairwise-compares distinct normalized names within same-first-word blocks.
    Pairs only — no transitive clustering. Ported from the Monitor app."""
    blocks: dict[str, list[str]] = defaultdict(list)
    for name in distinct_names:
        blocks[_fuzzy_block_key(name)].append(name)

    matches: list[tuple[str, str, int]] = []
    for names in blocks.values():
        if len(names) < 2 or len(names) > FUZZY_MAX_BLOCK_SIZE:
            continue
        for i in range(len(names)):
            for j in range(i + 1, len(names)):
                score = fuzz.token_sort_ratio(names[i], names[j])
                if score >= FUZZY_SIMILARITY_THRESHOLD:
                    matches.append((names[i], names[j], round(score)))
    return matches


UPSERT_SQL = (
    "INSERT INTO accounts_mirror (id, name, normalized_name, owner_id, owner_name, "
    "parent_account_id, phone, tax_id, identification_number, phone_digits, tax_id_digits, "
    "identification_number_digits, created_at_crm, synced_at) "
    "VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s) "
    "ON DUPLICATE KEY UPDATE name=VALUES(name), normalized_name=VALUES(normalized_name), "
    "owner_id=VALUES(owner_id), owner_name=VALUES(owner_name), "
    "parent_account_id=VALUES(parent_account_id), phone=VALUES(phone), tax_id=VALUES(tax_id), "
    "identification_number=VALUES(identification_number), phone_digits=VALUES(phone_digits), "
    "tax_id_digits=VALUES(tax_id_digits), identification_number_digits=VALUES(identification_number_digits), "
    "created_at_crm=VALUES(created_at_crm), synced_at=VALUES(synced_at)"
)
UPSERT_BATCH_SIZE = 1000  # this CRM tenant holds 125k+ Indonesia Accounts alone

FUZZY_UPSERT_SQL = (
    "INSERT INTO fuzzy_matches (normalized_name_a, normalized_name_b, similarity, computed_at) "
    "VALUES (%s, %s, %s, %s)"
)
FUZZY_INSERT_BATCH_SIZE = 1000


async def do_sync(trigger: str = "manual") -> dict:
    started_at = datetime.utcnow()
    seen = 0
    upserted = 0
    error = None
    run_id = None

    async with _pool.acquire() as conn, conn.cursor() as cur:
        await cur.execute(
            "INSERT INTO sync_runs (started_at, sync_trigger) VALUES (%s, %s)",
            (started_at, trigger),
        )
        await cur.execute("SELECT LAST_INSERT_ID()")
        (run_id,) = await cur.fetchone()

    try:
        accounts = await crm.crm_client.list_all_accounts(record_type_id=ACCOUNT_RECORD_TYPE_ID)
        seen = len(accounts)
        # Every row touched by this run gets the SAME synced_at, truncated to whole
        # seconds (MySQL DATETIME has no fractional precision) — see the Monitor app's
        # do_sync for why a microsecond-bearing "now" would otherwise delete every row
        # this run just inserted.
        now = datetime.utcnow().replace(microsecond=0)

        rows = [
            (
                acc["id"], acc["name"], normalize_name(acc["name"]), acc.get("owner_id"), acc.get("owner_name"),
                acc.get("parent_account_id"), acc.get("phone"), acc.get("tax_id"), acc.get("identification_number"),
                _normalize_identity_value(acc.get("phone")), _normalize_identity_value(acc.get("tax_id")),
                _normalize_identity_value(acc.get("identification_number")),
                _parse_dt(acc.get("created_at")), now,
            )
            for acc in accounts
            if acc.get("id") is not None and acc.get("name")
        ]

        async with _pool.acquire() as conn, conn.cursor() as cur:
            for i in range(0, len(rows), UPSERT_BATCH_SIZE):
                batch = rows[i : i + UPSERT_BATCH_SIZE]
                await cur.executemany(UPSERT_SQL, batch)
                upserted += len(batch)
                await cur.execute(
                    "UPDATE sync_runs SET accounts_seen=%s, upserted=%s WHERE id=%s",
                    (seen, upserted, run_id),
                )

            # Prune rows this run didn't touch, then recompute fuzzy matches against
            # the now-current name set — only when the pull actually returned
            # something, so a transient empty/error response never wipes the mirror.
            if upserted > 0:
                await cur.execute("DELETE FROM accounts_mirror WHERE synced_at < %s", (now,))

                await cur.execute("SELECT DISTINCT normalized_name FROM accounts_mirror")
                distinct_names = [r[0] for r in await cur.fetchall()]
                fuzzy_pairs = _compute_fuzzy_matches(distinct_names)
                fuzzy_rows = [(a, b, score, now) for a, b, score in fuzzy_pairs]
                for i in range(0, len(fuzzy_rows), FUZZY_INSERT_BATCH_SIZE):
                    await cur.executemany(FUZZY_UPSERT_SQL, fuzzy_rows[i : i + FUZZY_INSERT_BATCH_SIZE])
                await cur.execute("DELETE FROM fuzzy_matches WHERE computed_at < %s", (now,))
    except Exception as e:
        error = str(e)[:1000]

    finished_at = datetime.utcnow()
    async with _pool.acquire() as conn, conn.cursor() as cur:
        await cur.execute(
            "UPDATE sync_runs SET finished_at=%s, accounts_seen=%s, upserted=%s, error=%s WHERE id=%s",
            (finished_at, seen, upserted, error, run_id),
        )

    return {"accounts_seen": seen, "upserted": upserted, "error": error}


# ---------------------------------------------------------------------------
# ID resolution + candidate-group lookup
# ---------------------------------------------------------------------------

async def _fetch_mirror_by_id(conn, account_id: int) -> dict | None:
    async with conn.cursor() as cur:
        await cur.execute(
            f"SELECT {', '.join(_MIRROR_FIELDS)} FROM accounts_mirror WHERE id=%s", (account_id,)
        )
        row = await cur.fetchone()
    return _row_to_dict(row) if row else None


async def _fetch_mirror_rows_by_normalized_name(conn, normalized_name: str, exclude_id: int) -> list[dict]:
    if not normalized_name:
        return []
    async with conn.cursor() as cur:
        await cur.execute(
            f"SELECT {', '.join(_MIRROR_FIELDS)} FROM accounts_mirror WHERE normalized_name=%s AND id != %s",
            (normalized_name, exclude_id),
        )
        rows = await cur.fetchall()
    return [_row_to_dict(r) for r in rows]


async def _fetch_identity_group_members(conn, digits_col: str, value: str, exclude_id: int) -> list[dict]:
    # digits_col comes only from the fixed IDENTITY_FIELDS whitelist above, never from
    # request input, so interpolating the column name is safe.
    async with conn.cursor() as cur:
        await cur.execute(
            f"SELECT {', '.join(_MIRROR_FIELDS)} FROM accounts_mirror WHERE {digits_col}=%s AND id != %s",
            (value, exclude_id),
        )
        rows = await cur.fetchall()
    return [_row_to_dict(r) for r in rows]


async def _fetch_fuzzy_pairs(conn, normalized_name: str) -> list[tuple[str, int]]:
    async with conn.cursor() as cur:
        await cur.execute(
            "SELECT normalized_name_b, similarity FROM fuzzy_matches WHERE normalized_name_a=%s "
            "UNION ALL "
            "SELECT normalized_name_a, similarity FROM fuzzy_matches WHERE normalized_name_b=%s",
            (normalized_name, normalized_name),
        )
        rows = await cur.fetchall()
    return [(r[0], r[1]) for r in rows]


async def _resolve_names_batch(conn, ids: set) -> dict:
    """Batch-resolves Account id -> {id, name, owner_name}, first from the local
    mirror, then live from CRM for any ids the mirror doesn't have (e.g. a Parent
    Account that isn't itself an Indonesia Account, or was created after the last
    sync). Used to resolve current/suggested Parent Account names."""
    ids = {i for i in ids if i}
    if not ids:
        return {}
    placeholders = ",".join(["%s"] * len(ids))
    async with conn.cursor() as cur:
        await cur.execute(
            f"SELECT id, name, owner_name FROM accounts_mirror WHERE id IN ({placeholders})",
            tuple(ids),
        )
        rows = await cur.fetchall()
    result = {r[0]: {"id": r[0], "name": r[1], "owner_name": r[2]} for r in rows}
    missing = ids - result.keys()
    for mid in missing:
        live = await crm.crm_client.get_record("Account", mid)
        if live:
            result[mid] = {"id": mid, "name": live.get("name"), "owner_name": live.get("owner_name")}
    return result


def _suggest_parent(members: list[dict]) -> tuple[int, bool]:
    """Same logic as the Monitor app's _summarize_group: the existing parent used by
    the MOST members (majority vote, so a member that's itself a hub within the group
    doesn't lose to its own upstream parent), else the oldest-created member."""
    parent_counts: dict[int, int] = defaultdict(int)
    for m in members:
        if m["parent_account_id"]:
            parent_counts[m["parent_account_id"]] += 1
    if parent_counts:
        return max(parent_counts, key=lambda pid: parent_counts[pid]), True
    oldest = min(members, key=lambda m: (m["created_at_crm"] or datetime.min))
    return oldest["id"], False


async def _build_group(conn, match_type: str, matched_label: str, members: list[dict]) -> dict:
    suggested_parent_id, has_existing_parent = _suggest_parent(members)
    parent_ids = {m["parent_account_id"] for m in members if m["parent_account_id"]}
    parent_ids.add(suggested_parent_id)
    names = await _resolve_names_batch(conn, parent_ids)
    suggested_parent_name = names.get(suggested_parent_id, {}).get("name")

    rows = []
    for m in members:
        if m["id"] == suggested_parent_id:
            status = "is_suggested_parent"
        elif m["parent_account_id"] == suggested_parent_id:
            status = "already_under_suggested_parent"
        else:
            status = "needs_mapping"
        current_parent = names.get(m["parent_account_id"]) if m["parent_account_id"] else None
        rows.append({
            "id": m["id"],
            "name": m["name"],
            "owner_id": m["owner_id"],
            "owner_name": m["owner_name"],
            "current_parent_id": m["parent_account_id"],
            "current_parent_name": current_parent["name"] if current_parent else None,
            "status": status,
        })
    rows.sort(key=lambda r: r["name"])

    return {
        "match_type": match_type,
        "matched_label": matched_label,
        "has_existing_parent": has_existing_parent,
        "suggested_parent_id": suggested_parent_id,
        "suggested_parent_name": suggested_parent_name,
        "members": rows,
    }


async def _lookup_account_groups(conn, account: dict) -> list[dict]:
    """Finds this Account's candidate duplicate groups — exact name, fuzzy name, and
    shared phone/NPWP/identification number — the same three signals the Monitor app
    uses, but scoped to just this one Account instead of the whole mirror."""
    groups = []
    target_member = {f: account.get(f) for f in _MIRROR_FIELDS}

    exact_others = await _fetch_mirror_rows_by_normalized_name(conn, account["normalized_name"], account["id"])
    if exact_others:
        groups.append(await _build_group(conn, "exact_name", "Exact Name Match", [target_member] + exact_others))

    fuzzy_pairs = await _fetch_fuzzy_pairs(conn, account["normalized_name"])
    for other_name, similarity in fuzzy_pairs:
        others = await _fetch_mirror_rows_by_normalized_name(conn, other_name, account["id"])
        if others:
            group = await _build_group(
                conn, "fuzzy_name", f"Fuzzy Name Match ({similarity}%)", [target_member] + others
            )
            group["similarity"] = similarity
            groups.append(group)

    for digits_col, raw_col, label in IDENTITY_FIELDS:
        value = _normalize_identity_value(account.get(raw_col))
        if not value:
            continue
        others = await _fetch_identity_group_members(conn, digits_col, value, account["id"])
        if others:
            group = await _build_group(conn, "identity", f"Shared {label}", [target_member] + others)
            group["matched_field"] = label
            group["matched_value"] = account.get(raw_col)
            groups.append(group)

    return groups


def _account_from_live(record_id: int, acc: dict) -> dict:
    name = acc.get("name") or ""
    return {
        "id": record_id,
        "name": name,
        "normalized_name": normalize_name(name),
        "owner_id": acc.get("owner_id"),
        "owner_name": acc.get("owner_name"),
        "parent_account_id": acc.get("parent_account_id"),
        "phone": acc.get("phone"),
        "tax_id": acc.get("tax_id"),
        "identification_number": acc.get("identification_number"),
        "created_at_crm": _parse_dt(acc.get("created_at")),
    }


async def _resolve_account(conn, input_id: int) -> tuple[dict | None, bool, str | None, str | None, str | None]:
    """Resolves a raw input id to an Account row, trying (in order): the local mirror
    as an Account id; a live Opportunity lookup (-> its account_id) as a Fallback;
    then a live Account lookup as a last resort (covers an Account created after the
    last sync, or a wrong/edge-case id). Returns
    (account, in_mirror, note, input_type, opportunity_stage) where input_type is
    "account" or "opportunity" (which CRM object the input id turned out to be), or
    None when the id resolved to nothing — so the UI/export can show, for an
    Opportunity ID, the Account ID it maps to. opportunity_stage is that Opportunity's
    CRM `stage` field (e.g. "New", "Proposal Submitted") when input_type is
    "opportunity", else None — Account lookups have no single Opportunity to report a
    stage for.

    This is a best-effort type sniff, not a guarantee: Account and Opportunity ids are
    separate CRM sequences with non-overlapping observed ranges (confirmed live:
    Accounts ~1.4M vs Opportunities ~900K), so a numeric collision between the two
    object types is unlikely but not impossible — a known limitation, not a bug."""
    row = await _fetch_mirror_by_id(conn, input_id)
    if row:
        return row, True, None, "account", None

    opp = await crm.crm_client.get_record("Opportunity", input_id)
    if opp and opp.get("account_id"):
        acc_id = int(opp["account_id"])
        stage = opp.get("stage")
        row = await _fetch_mirror_by_id(conn, acc_id)
        if row:
            return row, True, None, "opportunity", stage
        acc = await crm.crm_client.get_record("Account", acc_id)
        if acc:
            return _account_from_live(acc_id, acc), False, "not_in_mirror", "opportunity", stage
        return None, False, "opportunity_account_not_found", None, None

    acc = await crm.crm_client.get_record("Account", input_id)
    if acc:
        return _account_from_live(input_id, acc), False, "not_in_mirror", "account", None

    return None, False, "not_found", None, None


async def _lookup_single(conn, input_id: int) -> dict:
    account, in_mirror, note, input_type, opportunity_stage = await _resolve_account(conn, input_id)
    if account is None:
        return {
            "input_id": input_id, "resolved": False, "note": note or "not_found",
            "input_type": input_type, "opportunity_stage": opportunity_stage,
            "account_id": None, "account_name": None, "account_owner": None,
            "current_parent_id": None, "current_parent_name": None,
            "in_mirror": False, "groups": [],
        }

    groups = await _lookup_account_groups(conn, account)
    current_parent_name = None
    if account.get("parent_account_id"):
        names = await _resolve_names_batch(conn, {account["parent_account_id"]})
        current_parent_name = names.get(account["parent_account_id"], {}).get("name")

    return {
        "input_id": input_id, "resolved": True, "note": note,
        "input_type": input_type, "opportunity_stage": opportunity_stage,
        "account_id": account["id"], "account_name": account["name"],
        "account_owner": account.get("owner_name"),
        "current_parent_id": account.get("parent_account_id"),
        "current_parent_name": current_parent_name,
        "in_mirror": in_mirror, "groups": groups,
    }


# Per-Account Opportunity-stage enrichment for the duplicate-match tables — single
# lookup only (GET /api/lookup), not batch runs: it's one extra live CRM call per
# distinct Account shown across a lookup's groups, and this CRM tenant has already
# been rate-limited (429) once today under much lighter load than a large batch would
# add, so this stays scoped to the one-ID-at-a-time request path.
_ACQUISITION_TYPE = "Acquisition"
# Confirmed live today: this CRM tenant 429'd repeatedly under lighter load than this
# (see the sync job's own MAX_CONCURRENT_PAGES comment in crm_client.py) — kept low.
_STAGE_ENRICH_CONCURRENCY = 2


async def _fetch_earliest_acquisition_opportunity(account_id: int) -> dict | None:
    """That Account's earliest-created Acquisition-type Opportunity (earliest by
    created_at) — chosen over "most recent" because Sales Ops wants to see progress
    from where an Account's acquisition effort actually started, not whatever
    Opportunity happens to be newest. Returns {"name", "stage"}, or None if the
    Account has no Acquisition Opportunity."""
    opps = await crm.crm_client.list_opportunities_by_account(account_id, opportunity_type=_ACQUISITION_TYPE)
    if not opps:
        return None
    earliest = min(opps, key=lambda o: _parse_dt(o.get("created_at")) or datetime.max)
    return {"name": earliest.get("name"), "stage": earliest.get("stage")}


async def _enrich_groups_with_member_stages(groups: list[dict]) -> None:
    """Best-effort: this enrichment is a nice-to-have on top of the actual duplicate
    lookup, not something worth failing the whole request over. This CRM tenant is
    prone to 429s (confirmed live, same as the sync job elsewhere in this app) — a
    single Account's Opportunity fetch failing must not 500 the entire /api/lookup
    response, so failures here are swallowed to "no stage available" per Account
    rather than propagated."""
    account_ids = {m["id"] for g in groups for m in g["members"]}
    semaphore = asyncio.Semaphore(_STAGE_ENRICH_CONCURRENCY)

    async def fetch(account_id: int) -> tuple[int, dict | None]:
        async with semaphore:
            try:
                return account_id, await _fetch_earliest_acquisition_opportunity(account_id)
            except Exception:
                return account_id, None

    opportunities = dict(await asyncio.gather(*[fetch(aid) for aid in account_ids]))
    for g in groups:
        for m in g["members"]:
            opp = opportunities.get(m["id"])
            m["stage"] = opp["stage"] if opp else None
            m["opportunity_name"] = opp["name"] if opp else None


# ---------------------------------------------------------------------------
# Batch jobs (paste or Excel upload)
# ---------------------------------------------------------------------------

def _dedupe_ids(raw_values) -> list[int]:
    ids = []
    seen = set()
    for raw in raw_values:
        if raw is None:
            continue
        try:
            val = int(float(str(raw).strip().strip(",;")))
        except (ValueError, TypeError):
            continue
        if val not in seen:
            seen.add(val)
            ids.append(val)
    return ids


def _parse_ids_from_xlsx(content: bytes) -> list[int]:
    """Reads the first column of the first sheet, one id per row — mixed Account ID /
    Opportunity ID, auto-detected later during lookup. Non-numeric rows (e.g. a header
    like "Account ID") are silently skipped rather than failing the whole upload."""
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    values = (row[0] for row in ws.iter_rows(values_only=True) if row)
    return _dedupe_ids(values)


def _parse_ids_from_csv(content: bytes) -> list[int]:
    """Reads the first column of a CSV/TSV/plain-text file, one id per row — same
    contract as _parse_ids_from_xlsx (mixed IDs, header rows silently skipped). The
    delimiter is sniffed (comma/semicolon/tab/pipe) because Excel "Save as CSV" in
    some locales writes semicolons; decoding tolerates a UTF-8 BOM."""
    text = content.decode("utf-8-sig", errors="replace")
    delimiter = ","
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
    except csv.Error:
        pass  # single-column file with no delimiter — comma reader still yields col 0
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    values = (row[0] for row in reader if row)
    return _dedupe_ids(values)


async def _create_job(source: str, total: int) -> int:
    async with _pool.acquire() as conn, conn.cursor() as cur:
        await cur.execute(
            "INSERT INTO lookup_jobs (status, source, created_at, total) VALUES ('running', %s, %s, %s)",
            (source, datetime.utcnow(), total),
        )
        await cur.execute("SELECT LAST_INSERT_ID()")
        (job_id,) = await cur.fetchone()
    return job_id


async def _run_batch_job(job_id: int, ids: list[int]) -> None:
    results = []
    try:
        async with _pool.acquire() as conn:
            for i, input_id in enumerate(ids):
                results.append(await _lookup_single(conn, input_id))
                if (i + 1) % 10 == 0 or i + 1 == len(ids):
                    async with conn.cursor() as cur:
                        await cur.execute(
                            "UPDATE lookup_jobs SET processed=%s WHERE id=%s", (i + 1, job_id)
                        )
        async with _pool.acquire() as conn, conn.cursor() as cur:
            await cur.execute(
                "UPDATE lookup_jobs SET status='done', finished_at=%s, processed=%s, results_json=%s WHERE id=%s",
                (datetime.utcnow(), len(ids), json.dumps(results), job_id),
            )
    except Exception as e:
        async with _pool.acquire() as conn, conn.cursor() as cur:
            await cur.execute(
                "UPDATE lookup_jobs SET status='error', finished_at=%s, error=%s WHERE id=%s",
                (datetime.utcnow(), str(e)[:1000], job_id),
            )


EXPORT_COLUMNS = [
    "Input ID", "Input Type", "Opportunity Stage", "Account ID", "Account Name", "Account Owner",
    "Current Parent ID", "Current Parent Name", "Match Type", "Suggested Parent ID",
    "Suggested Parent Name", "Status",
]

_INPUT_TYPE_LABEL = {"account": "Account", "opportunity": "Opportunity"}


def _input_type_label(input_type) -> str:
    return _INPUT_TYPE_LABEL.get(input_type, "")


def _export_rows(results: list[dict]) -> list[list]:
    rows = []
    for r in results:
        input_type = _input_type_label(r.get("input_type"))
        stage = r.get("opportunity_stage")
        if not r["resolved"]:
            rows.append([r["input_id"], input_type, stage, None, "ID not found", None, None, None, "Error", None, None, "not_found"])
            continue
        if not r["groups"]:
            rows.append([
                r["input_id"], input_type, stage, r["account_id"], r["account_name"], r["account_owner"],
                r["current_parent_id"], r["current_parent_name"], "No duplicates found",
                None, None, "no_duplicates",
            ])
            continue
        for group in r["groups"]:
            for member in group["members"]:
                rows.append([
                    r["input_id"], input_type, stage, member["id"], member["name"], member["owner_name"],
                    member["current_parent_id"], member["current_parent_name"],
                    group["matched_label"], group["suggested_parent_id"], group["suggested_parent_name"],
                    member["status"],
                ])
    return rows


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------

class Health(BaseModel):
    status: str


class Me(BaseModel):
    email: str | None
    user: str | None


class SyncResult(BaseModel):
    accounts_seen: int
    upserted: int
    error: str | None = None


class SyncStartResult(BaseModel):
    started: bool
    already_running: bool
    error: str | None = None


class SyncStatus(BaseModel):
    in_progress: bool
    last_synced_at: str | None
    last_result: SyncResult | None
    current_run_started_at: str | None = None


class BatchIdsPayload(BaseModel):
    ids: list[str]


class BatchStartResult(BaseModel):
    job_id: int
    total: int


class BatchStatus(BaseModel):
    status: str
    total: int
    processed: int
    error: str | None = None


@app.get("/health", response_model=Health)
def health():
    return {"status": "ok"}


@app.get("/api/me", response_model=Me)
def me(request: Request):
    return {
        "email": request.headers.get("x-forwarded-email"),
        "user": request.headers.get("x-forwarded-user"),
    }


@app.post("/api/sync/crm", response_model=SyncStartResult)
async def sync_crm():
    """Starts a mirror sync in the background and returns immediately — a full
    paginated Account pull is too slow to fit inside one HTTP request/response cycle.
    Poll GET /api/sync/status for progress."""
    global _sync_in_progress
    if _pool is None:
        return {"started": False, "already_running": False, "error": "DATABASE_URL not set"}
    if _sync_in_progress:
        return {"started": False, "already_running": True}
    _sync_in_progress = True
    asyncio.create_task(_run_sync("manual"))
    return {"started": True, "already_running": False}


@app.get("/api/sync/status", response_model=SyncStatus)
async def sync_status():
    if _pool is None:
        return {"in_progress": False, "last_synced_at": None, "last_result": None}
    async with _pool.acquire() as conn, conn.cursor() as cur:
        # started_at is only meaningful while a run is in progress — it's the run's
        # own start time, so the UI can show "started N minutes ago" instead of
        # leaving the user guessing whether a long-running sync is still healthy.
        await cur.execute(
            "SELECT accounts_seen, upserted, error, started_at FROM sync_runs ORDER BY id DESC LIMIT 1"
        )
        row = await cur.fetchone()
        await cur.execute("SELECT MAX(synced_at) FROM accounts_mirror")
        (last_synced,) = await cur.fetchone()
    last_result = {"accounts_seen": row[0], "upserted": row[1], "error": row[2]} if row else None
    return {
        "in_progress": _sync_in_progress,
        "last_synced_at": str(last_synced) if last_synced else None,
        "last_result": last_result,
        "current_run_started_at": str(row[3]) if row and _sync_in_progress else None,
    }


@app.get("/api/lookup")
async def lookup(id: int):
    """Synchronous single-ID lookup — resolves the id (Account or Opportunity) and
    returns its candidate duplicate groups, each member enriched with its earliest
    Acquisition-Opportunity Stage (one live CRM call per distinct Account across all
    groups — see _enrich_groups_with_member_stages for why this is single-lookup-only)."""
    if _pool is None:
        raise HTTPException(503, "Database not configured")
    async with _pool.acquire() as conn:
        result = await _lookup_single(conn, id)
    if result["resolved"] and result["groups"]:
        await _enrich_groups_with_member_stages(result["groups"])
    return result


@app.post("/api/lookup/batch", response_model=BatchStartResult)
async def lookup_batch(payload: BatchIdsPayload):
    if _pool is None:
        raise HTTPException(503, "Database not configured")
    ids = _dedupe_ids(payload.ids)
    if not ids:
        raise HTTPException(400, "No valid IDs provided")
    if len(ids) > BATCH_MAX_IDS:
        raise HTTPException(400, f"Too many IDs — max {BATCH_MAX_IDS} per batch")
    job_id = await _create_job("paste", len(ids))
    asyncio.create_task(_run_batch_job(job_id, ids))
    return {"job_id": job_id, "total": len(ids)}


@app.post("/api/upload", response_model=BatchStartResult)
async def upload(file: UploadFile = File(...)):
    if _pool is None:
        raise HTTPException(503, "Database not configured")
    content = await file.read()
    filename = (file.filename or "").lower()
    try:
        if filename.endswith(".xlsx"):
            ids = _parse_ids_from_xlsx(content)
        elif filename.endswith((".csv", ".tsv", ".txt")):
            ids = _parse_ids_from_csv(content)
        else:
            # Unknown/blank extension — try Excel first, then fall back to CSV parsing.
            try:
                ids = _parse_ids_from_xlsx(content)
            except Exception:
                ids = _parse_ids_from_csv(content)
    except Exception:
        raise HTTPException(400, "Could not read the uploaded file — upload an .xlsx or .csv with IDs in the first column")
    if not ids:
        raise HTTPException(400, "No valid IDs found in the first column of the uploaded file")
    if len(ids) > BATCH_MAX_IDS:
        raise HTTPException(400, f"Too many IDs — max {BATCH_MAX_IDS} per batch")
    job_id = await _create_job("upload", len(ids))
    asyncio.create_task(_run_batch_job(job_id, ids))
    return {"job_id": job_id, "total": len(ids)}


@app.get("/api/lookup/batch/{job_id}/status", response_model=BatchStatus)
async def batch_status(job_id: int):
    if _pool is None:
        raise HTTPException(503, "Database not configured")
    async with _pool.acquire() as conn, conn.cursor() as cur:
        await cur.execute("SELECT status, total, processed, error FROM lookup_jobs WHERE id=%s", (job_id,))
        row = await cur.fetchone()
    if not row:
        raise HTTPException(404, "Job not found")
    return {"status": row[0], "total": row[1], "processed": row[2], "error": row[3]}


@app.get("/api/lookup/batch/{job_id}/results")
async def batch_results(job_id: int):
    if _pool is None:
        raise HTTPException(503, "Database not configured")
    async with _pool.acquire() as conn, conn.cursor() as cur:
        await cur.execute("SELECT status, results_json, error FROM lookup_jobs WHERE id=%s", (job_id,))
        row = await cur.fetchone()
    if not row:
        raise HTTPException(404, "Job not found")
    status, results_json, error = row
    if status == "error":
        raise HTTPException(500, error or "Batch job failed")
    if status != "done":
        raise HTTPException(409, "Batch job still running")
    return {"results": json.loads(results_json) if results_json else []}


@app.get("/api/lookup/batch/{job_id}/results.xlsx")
async def batch_results_xlsx(job_id: int):
    if _pool is None:
        raise HTTPException(503, "Database not configured")
    async with _pool.acquire() as conn, conn.cursor() as cur:
        await cur.execute("SELECT status, results_json FROM lookup_jobs WHERE id=%s", (job_id,))
        row = await cur.fetchone()
    if not row or row[0] != "done":
        raise HTTPException(404, "Results not available")
    results = json.loads(row[1]) if row[1] else []

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Parent Account Lookup"
    ws.append(EXPORT_COLUMNS)
    for row_values in _export_rows(results):
        ws.append(row_values)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=parent_account_lookup_results.xlsx"},
    )
